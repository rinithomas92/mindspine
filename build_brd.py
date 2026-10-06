from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path

OUT=Path('deliverables'); OUT.mkdir(exist_ok=True)
d=Document(); s=d.sections[0]
s.page_height=Inches(11.7); s.page_width=Inches(8.3)
s.top_margin=Inches(.7); s.bottom_margin=Inches(.65); s.left_margin=Inches(.75); s.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
st=d.styles['Normal']; st.font.size=Pt(10.5); st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.08
for name,size in [('Title',29),('Subtitle',14),('Heading 1',19),('Heading 2',12)]:
 d.styles[name].font.size=Pt(size)
 d.styles[name].paragraph_format.space_after=Pt(8)
d.core_properties.title='MindSpine Business Requirements Document'
d.core_properties.subject='Appointment booking and care management platform'
d.core_properties.author='MindSpine'
f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('MindSpine BRD  |  v0.1  |  ').font.size=Pt(9)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
def p(t): return d.add_paragraph(t)
def h(t): d.add_heading(t,2)
def page(t): d.add_page_break(); d.add_heading(t,1)
def req(i,title,text):
 pp=d.add_paragraph(); pp.add_run(i+'  '+title+'. ').bold=True; pp.add_run(text)
def table(headers,rows,widths):
 t=d.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Inches(w)
 for c,txt in zip(t.rows[0].cells,headers): c.text=txt
 for row in rows:
  for c,txt in zip(t.add_row().cells,row): c.text=txt
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr(); ns=OxmlElement('w:cantSplit'); pr.append(ns)
  if ri==0: pr.append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=c._tc.get_or_add_tcPr(); shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DCE8F4' if ri==0 else ('F5F7FA' if ri%2==0 else 'FFFFFF')); cp.append(shade)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    b=OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
   cp.append(borders); mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); mar.append(e)
   cp.append(mar)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(2); pp.paragraph_format.space_before=Pt(2)
    for r in pp.runs: r.font.size=Pt(10); r.bold=(ri==0)
 return t

d.add_heading('MindSpine Business Requirements Document',0)
d.add_paragraph('Appointment booking and care management platform',style='Subtitle')
p('Version 0.1  |  20 September 2026  |  For stakeholder review')
h('1 Purpose and executive overview')
p('MindSpine will provide a connected digital service for patients, chiropractors, and administrators to manage appointments, care records, payments, and reports. This BRD defines the business capabilities and acceptance conditions required to move from the system architecture to an agreed delivery scope.')
p('Stakeholders should approve the core requirements and resolve the decisions in Section 10 before delivery estimates and release commitments are finalized. This document is a requirements baseline for review; it does not record approval or evidence that the capabilities have been implemented.')
h('2 Business objectives')
table(['Objective','Intended business outcome','Suggested measure'],[
('BO1 Simplify booking','Patients can book and manage appointments digitally.','Booking completion rate and time to confirm'),
('BO2 Support continuity of care','Practitioners can review history and maintain treatment records.','Visits with complete clinical documentation'),
('BO3 Improve payment visibility','Bills and payment states remain linked to patient activity.','Payment success and reconciliation exceptions'),
('BO4 Improve communication','Patients receive timely reminders and report alerts.','Delivery rate and appointment no show rate'),
('BO5 Enable operational oversight','Administrators manage users, schedules, and analytics.','Schedule utilization and report accuracy')],[1.55,2.65,2.6])
p('Baseline values, target values, measurement windows, and accountable owners for these measures require business approval.')
h('Document basis')
p('Reference S1 is the supplied MindSpine System Architecture image, titled “Appointment Booking • Care Management • Payments • Reports”, file ChatGPT Image Sep 20, 2026, 12_23_59 PM.png. Capability groups trace to S1 in Section 9. Detailed controls, workflow rules, acceptance conditions, and release priorities below are proposed elaborations for approval.')

page('3 Stakeholders and scope')
h('User roles')
table(['Role','Business responsibilities','Proposed access boundary'],[
('Patient or user','Register, book, view history and reports, pay, and receive notifications.','Own profile, appointments, bills, and released records only.'),
('Chiropractor','Manage bookings, review patient history, document care, and export authorized data.','Patients within an approved care relationship; sharing rules require confirmation.'),
('Administrator','Manage users, schedules, analytics, and configuration.','Operational access by permission; clinical access must be explicitly authorized.')],[1.15,2.8,2.85])
p('Business sponsor, clinical lead, operations lead, finance representative, privacy or security owner, and delivery lead are proposed review roles. Named owners remain to be assigned.')
h('Proposed core release')
p('The core release includes registration and role based login; patient profiles and history; appointment booking, rescheduling, and cancellation; practitioner notes, diagnoses, treatment plans, and uploads; bills, invoices, online payments and payment status; reminders and alerts; patient reports with download and print; practitioner Excel exports; and administration of users, schedules, analytics, and settings.')
p('Patient access must support web and mobile use. The image lists responsive web, iOS, and Android, but does not establish whether separate native applications are required. Delivery format and supported devices are an approval decision.')
h('Integration scope')
p('Payment, email, and SMS services support the core journeys. Push notifications, WhatsApp, and calendar synchronization are in the target capability scope; their release timing depends on channel selection, consent rules, provider access, and platform delivery decisions.')
h('Future scope and exclusions')
p('Laboratory, insurance, and external electronic medical record integrations are explicitly future expansion in S1. Teleconsultation, pharmacy or inventory management, insurance claims processing, automated clinical diagnosis, and historical data migration are not included in the proposed core baseline unless separately approved.')
h('Priority convention')
p('BR01 through BR24 are proposed core requirements. Calendar synchronization is a target integration with timing to confirm. Security and operational requirements are proposed release conditions. These priorities express a recommended baseline, not an approved delivery commitment.')

page('4 Identity patient and appointment requirements')
req('BR01','Registration and authentication','The platform shall support patient registration, login for all three roles, and password reset. Practitioner and administrator provisioning shall follow an approved onboarding process. Acceptance: an active user can sign in and reset access; invalid credentials and disabled accounts are rejected.')
req('BR02','Role based authorization','The platform shall enforce access permissions for screens, APIs, records, downloads, and exports. Acceptance: role and ownership tests prevent a patient from accessing another patient’s data and prevent unauthorized role elevation.')
req('BR03','Patient profile','The platform shall maintain patient identity, contact details, and relevant profile information. Acceptance: authorized profile changes persist and are reflected in the patient record; required fields and duplicate handling are agreed before configuration.')
req('BR04','Patient history and documents','The platform shall maintain medical history, visit history, and patient documents. Acceptance: authorized users can retrieve the correct patient’s history, while the patient sees only records approved for patient access.')
req('BR05','Schedule and slot management','Authorized users shall configure practitioner availability and appointment slots. Acceptance: only eligible slots are offered; unavailable periods cannot be booked. Appointment durations, time zones, buffers, and clinic structure require confirmation.')
req('BR06','Online booking','Patients shall select an available slot and submit a booking linked to their profile and practitioner. Acceptance: a confirmed booking has a unique reference; concurrent attempts for the same exclusive slot cannot both succeed.')
req('BR07','Reschedule and cancellation','Patients and authorized staff shall reschedule or cancel within agreed policy. Acceptance: an allowed change updates availability and appointment history and triggers the relevant notification; a prohibited change explains the applicable restriction.')
req('BR08','Practitioner booking management','Chiropractors shall view and manage their bookings within assigned permissions. Acceptance: schedule changes appear consistently in the practitioner and patient views, and one practitioner cannot alter another’s bookings without permission.')
h('Booking business rules requiring approval')
p('Define when a booking becomes confirmed, whether payment is required first, and how long a slot may be held. Approve cancellation cutoffs, rescheduling limits, no show handling, and refund eligibility. These rules must be configurable or documented before user acceptance testing.')

page('5 Clinical and financial requirements')
req('BR09','Clinical documentation','Chiropractors shall create and update visit notes, diagnoses, treatments, and treatment plans for authorized patients. Acceptance: each entry is linked to the patient, visit where applicable, practitioner, and date; changes remain attributable to the editor.')
req('BR10','Clinical file uploads','Authorized practitioners shall upload and retrieve supporting files such as X rays and reports. Acceptance: files remain linked to the correct patient and are inaccessible to unauthorized users. Allowed formats, size limits, and patient upload permissions require approval.')
req('BR11','Patient record visibility','Patients shall view their available visit history and published clinical reports. Acceptance: a released report is visible to the correct patient and inaccessible to other patients. Note visibility and the report release process must be approved by the clinical lead.')
req('BR12','Bills and invoices','The platform shall generate patient bills and invoices for applicable appointments or services. Acceptance: each document has a unique reference, patient association, itemized amounts, total, and payment state. Currency, taxes, numbering, and templates require finance approval.')
req('BR13','Online payments','Patients shall make payments through a selected payment gateway. Acceptance: successful, failed, pending, and abandoned payment attempts receive an accurate status; a return to the application alone does not establish payment success.')
req('BR14','Payment confirmation and reconciliation','Payment updates shall be verified and linked to the correct bill and booking. Acceptance: duplicate provider events do not create duplicate receipts or charges in the platform; delayed confirmations can be reconciled and unresolved exceptions are visible to authorized staff.')
req('BR15','Refund management','Authorized staff shall initiate or record eligible refunds through the payment integration. Acceptance: refund status, amount, original payment reference, and actor are recorded; duplicate processing is prevented and failed refunds are available for follow up.')
h('Proposed integrity rules')
p('Appointments and financial transactions must have separate status histories. A payment failure must not silently confirm a booking that requires payment. A cancellation must not imply a completed refund. Corrected clinical and financial information must retain sufficient history for authorized review.')
p('The specific status model, clinical amendment process, refund approval limits, partial payment support, and partial refund support require business decisions. No clinical recommendations or automated diagnoses are part of these requirements.')

page('6 Communications reporting and administration')
req('BR16','Notifications and reminders','The platform shall trigger booking and change notifications, appointment reminders, payment updates, and report ready alerts. Acceptance: each enabled event sends the approved template to the correct recipient through the configured channel without unnecessary duplicate messages.')
req('BR17','Channel preferences and delivery tracking','The platform shall support the approved email, SMS, push, and optional WhatsApp channels. Acceptance: channel preferences and applicable consent choices are respected; delivery failures are recorded and retried under an agreed policy. Reminder lead times and fallback rules require approval.')
req('BR18','Patient reports and printing','Authorized users shall generate and share patient bills and clinical or diagnosis reports. Acceptance: approved reports can be downloaded in PDF and printed with correct patient identity, dates, and content; patients can retrieve reports released to them.')
req('BR19','Excel export','Chiropractors shall export authorized patient data and reports to Excel. Acceptance: the exported dataset matches the selected scope and approved field list, excludes inaccessible records, and has usable column headings and dates.')
req('BR20','User administration','Administrators shall create or manage users, roles, and account status. Acceptance: approved changes take effect across application access, and deactivation prevents further access without deleting required records.')
req('BR21','Schedule administration','Administrators shall manage practitioner schedules and availability. Acceptance: a schedule change does not silently remove an existing booking; affected appointments are identified for authorized resolution and notification.')
req('BR22','Analytics','Administrators shall view operational analytics. Proposed measures include booking volume, appointment status, utilization, billed amounts, collections, and refunds. Acceptance: agreed date filters and definitions reconcile to the underlying authorized records.')
req('BR23','Configuration','Administrators shall maintain approved system settings and business parameters. Acceptance: only authorized administrators can modify settings; changes to patient facing rules take effect predictably and are traceable.')
req('BR24','Report access and sharing','Reports, exports, and shared files shall retain the same access restrictions as their source records. Acceptance: an unauthorized user cannot retrieve a file using a copied identifier or link; report sharing and download activity can be reviewed by authorized staff.')
h('Communication rule')
p('Notification templates should use only the information needed for the event and direct patients to authenticated access for sensitive details. Template content, consent handling, and permitted channels require stakeholder approval.')

page('7 Data integrations and solution boundaries')
h('Business data')
table(['Entity','Minimum business information','Proposed accountable role'],[
('User and patient','Identity, role, account state, contact information, preferences, and history','Operations with clinical oversight'),
('Practitioner and schedule','Practitioner identity, availability, slots, and exceptions','Operations'),
('Appointment and visit','Patient, practitioner, time, status, change history, and visit association','Operations and clinical lead'),
('Clinical record and file','Notes, diagnosis, treatment plan, author, timestamps, and document association','Clinical lead'),
('Bill payment and refund','Amounts, currency, status, references, and reconciliation history','Finance'),
('Report notification and audit','Output scope, recipient, event, delivery state, actor, and timestamp','Operations and security owner')],[1.45,3.5,1.85])
h('External dependencies')
p('Payment gateway: select Razorpay, Stripe, PayU, or another approved provider. Required outcomes are online payment, refund handling, and verified status updates through webhooks or equivalent provider interfaces.')
p('Messaging: select email and SMS providers; S1 offers SendGrid or AWS SES and Twilio or MSG91 as examples. WhatsApp and push require confirmed channels, credentials, message approvals where applicable, and appropriate client support.')
p('Calendar integration: synchronize appointments with Google Calendar or Outlook if included in the release. Approve synchronization direction, conflict handling, update latency, and the minimum patient information exposed. A cancelled or rescheduled appointment must not leave a misleading active event.')
h('Architecture interpretation')
p('S1 separates patient, practitioner, and administrator front ends from an API gateway and seven service domains: authentication, patients, appointments, clinical records, billing and payments, notifications, and reports. These boundaries describe the target solution; they do not require separate deployments in the BRD.')
p('The image proposes HTTPS, authentication, rate limiting, and load balancing, supported by a relational database, file storage, cache, and monitoring. PostgreSQL or MySQL; Supabase, AWS S3, or Azure Blob; Redis; and Prometheus with Grafana are candidate technologies, not final procurement decisions. External integrations are logical service connections; the diagram does not authorize direct external access to the primary database.')

page('8 Quality and operational requirements')
p('The following controls and numerical targets are proposed for approval. They elaborate the secure and scalable architecture objective and must be validated against business volume, hosting choices, budget, and applicable obligations.')
table(['ID','Requirement and acceptance measure'],[
('NFR01 Security','Use encrypted transport, protect stored sensitive data, secure credentials, and enforce authorization server side. Acceptance: approved security testing finds no unresolved critical or high risk issues before release.'),
('NFR02 Privacy and access','Agree data collection, permitted use, patient consent where required, retention, deletion, and residency rules. Acceptance: approved rules are implemented and unauthorized access tests pass.'),
('NFR03 Auditability','Record access and material changes to clinical records, permissions, financial states, and exports with actor and time. Acceptance: authorized reviewers can reconstruct representative events without exposing secrets in logs.'),
('NFR04 Performance','Proposed target: 95 percent of routine interactive actions complete within 3 seconds, excluding external provider delays. Confirm load profile and dataset before testing; measure provider delay separately.'),
('NFR05 Availability','Proposed target: 99.5 percent monthly service availability, with planned maintenance treatment agreed. Acceptance: monitoring can measure availability and alert on failures.'),
('NFR06 Recovery','Proposed targets: recovery point within 24 hours and recovery time within 8 hours. Acceptance: a restoration exercise proves recoverability of records and associated files; approve backup retention separately.'),
('NFR07 Usability','Support agreed desktop and mobile screen sizes and accessible core journeys. Acceptance: users complete booking, payment, and report access using the approved browser and device matrix, including keyboard navigation.'),
('NFR08 Reliability and scale','Prevent duplicate bookings and financial effects; support controlled retries and capacity growth. Acceptance: concurrent booking, duplicate webhook, and provider outage scenarios preserve record consistency.'),
('NFR09 Operations','Monitor errors, service health, failed integrations, and delivery queues. Acceptance: an induced failure raises an actionable alert, and an assigned support owner can follow a documented recovery procedure.')],[1.2,5.6])
h('Controls requiring local determination')
p('Applicable jurisdiction, healthcare and privacy obligations, clinical record retention, tax rules, and payment responsibilities must be established by the responsible business owners. No specific legal compliance certification is asserted in this BRD.')

page('9 Business workflows and acceptance')
h('End to end workflows')
p('Booking and payment: patient signs in, selects an available slot, reviews the charge, and submits the booking. Where advance payment is required, the slot follows an agreed hold policy until payment verification. Confirmation updates both schedules and triggers notification. Failure or timeout follows the approved release and retry rules.')
p('Care and report delivery: practitioner opens an authorized patient record, reviews history, records the visit and treatment plan, and uploads supporting documents. The approved report release process makes the report available to the patient and triggers an alert. The patient downloads or prints the report.')
p('Appointment change and refund: patient or authorized staff requests a change; the platform checks policy, updates the appointment and availability, and sends a notification. Any refund follows a separate eligibility and authorization process, with its status tracked to completion or exception.')
h('Source traceability')
table(['S1 capability group','Requirements','Key acceptance evidence'],[
('Login and role based users','BR01–BR02, BR20','Role tests and account lifecycle checks'),
('Patient history and clinical care','BR03–BR04, BR09–BR11','Authorized visit and record lifecycle'),
('Booking and schedules','BR05–BR08, BR21','Booking concurrency and change tests'),
('Billing and payment','BR12–BR15','Success, failure, duplicate event, and refund tests'),
('Notifications','BR16–BR17','Event, preference, and delivery failure tests'),
('Reports and Excel download','BR18–BR19, BR24','PDF, print, Excel, and access checks'),
('Administration and architecture','BR22–BR23, NFR01–NFR09','Analytics reconciliation and operational tests')],[2.15,1.8,2.85])
h('Release acceptance conditions')
p('Business representatives must accept all approved core requirements through user acceptance testing. Evidence must cover each role, normal journeys, access denials, booking collisions, failed and delayed payments, duplicate provider events, notification failures, report outputs, and backup restoration.')
p('Proposed release gate: no unresolved critical or high severity defects; all agreed quality targets met; production integrations and support procedures ready; training completed for administrators and practitioners; and named business, clinical, finance, and security owners approve release. Any exception requires a recorded owner and resolution plan.')

page('10 Assumptions decisions and approval')
h('Planning assumptions')
p('The initial baseline assumes individual patient accounts, appointment based care, and practitioner authored clinical content. Clinic count, practitioner count, patient volume, geographic reach, caregiver access, and existing data migration needs remain to be confirmed. No schedule, cost, vendor contract, or implementation completion is implied.')
h('Decisions required before delivery commitment')
table(['ID','Decision','Proposed decision owner'],[
('D01','Clinic structure, service types, appointment durations, locations, time zones, and expected volumes','Business and operations'),
('D02','Responsive web versus native iOS and Android apps; supported devices and first release channels','Product owner'),
('D03','Payment timing, slot hold expiry, cancellation policy, no shows, refunds, taxes, and currency','Operations and finance'),
('D04','Practitioner access scope, patient note visibility, report release, amendments, and required clinical fields','Clinical lead'),
('D05','Jurisdiction, consent, retention, residency, caregiver or minor access, and audit retention','Privacy and clinical owners'),
('D06','Provider selection, calendar direction, credentials, costs, channel consent, and integration release timing','Product and technical leads'),
('D07','Performance load, availability, recovery targets, support coverage, and incident ownership','Business and technical leads'),
('D08','Analytics definitions, export fields, report templates, migration, budget, delivery phases, and timeline','Business sponsor')],[.55,4.7,1.55])
h('Principal risks and dependencies')
p('Unresolved booking and refund policies may cause inconsistent patient outcomes; resolve D03 before implementation. Broad record or export access may expose sensitive information; approve D04 and D05 and verify access tests. Provider delays or outages may disrupt payments and reminders; agree exception handling and reconciliation before launch. Unconfirmed native app or migration scope may alter cost and timing; close D02 and D08 before estimation.')
h('Approval record')
p('Approval is pending. Record the approver name, role, date, document version, decision, and conditions for the business sponsor, clinical lead, operations and finance representatives, privacy or security owner, and technical lead. Changes after approval must identify impacted requirements and receive the appropriate owner’s approval.')

for root in [d.styles.element, d.element]:
 for border in list(root.iter(qn('w:pBdr'))): border.getparent().remove(border)
d.save(OUT/'MindSpine_Business_Requirements_Document.docx')
print(str((OUT/'MindSpine_Business_Requirements_Document.docx').resolve()))
